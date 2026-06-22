from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = r"C:\Users\juan.cornejo\Documents\gugnir v2\Cobertura_IoT_Austral_Polar_Brief.docx"


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="DADCE0", size=8):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def apply_font(run, size, bold=False, color="000000", name="Arial"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.orientation = WD_ORIENTATION.PORTRAIT
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(11)

for style_name, size, color, before, after in (
    ("Heading 1", 20, "000000", 20, 6),
    ("Heading 2", 16, "000000", 18, 6),
    ("Heading 3", 14, "434343", 16, 4),
):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = False
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(3)
run = title.add_run("Cobertura Satelital IoT en Zonas Australes y Polares")
apply_font(run, 26, bold=False)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(10)
subtitle_run = subtitle.add_run(
    "Brief de trabajo para la generación de mapas de calor de cobertura comparativa por constelación"
)
apply_font(subtitle_run, 11, color="555555")

meta_table = doc.add_table(rows=6, cols=2)
meta_table.alignment = WD_TABLE_ALIGNMENT.LEFT
meta_table.autofit = False
meta_table.columns[0].width = Inches(1.9)
meta_table.columns[1].width = Inches(4.6)
set_table_borders(meta_table, color="DADCE0", size=8)

metadata = [
    (
        "Asignación",
        "Cobertura satelital comparativa para comunicaciones IoT en latitudes australes y polares.",
    ),
    (
        "Sistema de Impacto",
        "Motor de análisis geoespacial, simulador de conectividad y tablero operativo de seguimiento satelital.",
    ),
    (
        "Objetivo",
        "Generar mapas de calor que permitan comparar la cobertura efectiva de distintas constelaciones satelitales orientadas a IoT, con foco en zonas australes y polares, identificando bandas de frecuencia, número de satélites activos, huella de cobertura, velocidades de transmisión y disponibilidad de hardware terminal.",
    ),
    (
        "Entregable Requerido",
        "Construir un set de productos técnicos compuesto por capas geoespaciales, mapas rasterizados de intensidad de cobertura y una matriz estructurada en CSV o GeoJSON. Cada registro debe incluir constelación, marca de tiempo, coordenadas, nivel estimado de cobertura, frecuencia operativa, elevación, throughput esperado, visibilidad y estado de disponibilidad de terminales compatibles.",
    ),
    (
        "Foco de Integracion",
        "El resultado debe alimentar modelos de planificación operativa, simulaciones de continuidad de servicio y evaluaciones logísticas para despliegues IoT en entornos extremos.",
    ),
    (
        "Resultado Esperado",
        "Una línea base comparativa que permita priorizar constelaciones y equipos por región, perfil de tráfico y resiliencia operacional.",
    ),
]

for row, (label, value) in zip(meta_table.rows, metadata):
    row.cells[0].width = Inches(1.9)
    row.cells[1].width = Inches(4.6)
    for idx, text in enumerate((label, value)):
        cell = row.cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        apply_font(r, 10.5, bold=(idx == 0), color="000000" if idx == 0 else "111111")

doc.add_paragraph("")

h1 = doc.add_paragraph("Estructura de la Presentación Final", style="Heading 1")
h1.paragraph_format.space_after = Pt(6)

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(8)
r1 = intro.add_run("Formato sugerido: ")
apply_font(r1, 11, bold=True)
r2 = intro.add_run(
    "presentación ejecutiva de 15 minutos con soporte cartográfico, comparación por constelación y cierre de recomendación técnica."
)
apply_font(r2, 11)

sections = [
    (
        "1. El Problema Operativo (2 minutos)",
        [
            "Explicar por qué las latitudes australes y polares concentran brechas de visibilidad, menor ángulo de elevación y mayores riesgos de intermitencia para enlaces IoT.",
            "Precisar la necesidad de una referencia cuantitativa para estimar continuidad de servicio y planificar despliegues remotos.",
        ],
    ),
    (
        "2. La Solución Desarrollada (3 minutos)",
        [
            "Describir la metodología para consolidar efemérides, parámetros orbitales, footprint estimado y capas geoespaciales de cobertura.",
            "Definir el criterio de comparación entre constelaciones: frecuencia, cantidad de satélites, área de cobertura, velocidad de transmisión y madurez del ecosistema de hardware.",
        ],
    ),
    (
        "3. Demostración Práctica (5 minutos)",
        [
            "Mostrar los mapas de calor por región y por constelación, destacando gradientes de cobertura, persistencia temporal y zonas de sombra.",
            "Explicar cómo leer la matriz estructurada y cómo vincular cada pixel o celda geográfica con disponibilidad estimada de servicio.",
        ],
    ),
    (
        "4. Recomendación Comparativa (3 minutos)",
        [
            "Sintetizar ventajas y limitaciones de cada constelación para telemetría de baja tasa, mensajería intermitente y sensorización ambiental.",
            "Priorizar alternativas según cobertura polar, ancho de banda requerido y disponibilidad comercial de terminales o módulos integrables.",
        ],
    ),
    (
        "5. Integración Operativa (2 minutos)",
        [
            "Definir cómo incorporar la capa de cobertura en motores de simulación, tableros de monitoreo y modelos de riesgo operacional.",
            "Proponer una hoja de ruta para actualización periódica de datos y validación en terreno.",
        ],
    ),
]

for heading, bullets in sections:
    p = doc.add_paragraph(heading, style="Heading 2")
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    for bullet in bullets:
        bp = doc.add_paragraph(style="List Bullet")
        bp.paragraph_format.space_after = Pt(4)
        bp.paragraph_format.line_spacing = 1.15
        br = bp.add_run(bullet)
        apply_font(br, 11)

doc.add_paragraph("Variables Mínimas por Constelación", style="Heading 2")

vars_table = doc.add_table(rows=1, cols=5)
vars_table.alignment = WD_TABLE_ALIGNMENT.LEFT
vars_table.autofit = False
for width, col in zip((1.3, 1.15, 1.3, 1.05, 1.7), vars_table.columns):
    col.width = Inches(width)
set_table_borders(vars_table, color="DADCE0", size=8)

headers = [
    "Variable",
    "Unidad",
    "Uso",
    "Prioridad",
    "Observacion",
]
for idx, text in enumerate(headers):
    cell = vars_table.rows[0].cells[idx]
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    r = p.add_run(text)
    apply_font(r, 10, bold=True)

rows = [
    ("Frecuencia", "MHz / GHz", "Compatibilidad regulatoria y de enlace", "Alta", "Separar uplink y downlink si aplica."),
    ("Número de satélites", "conteo", "Densidad orbital y revisita", "Alta", "Distinguir activos y planificados."),
    ("Área de cobertura", "km2 / huella", "Alcance geográfico real", "Alta", "Incluir foco polar y marítimo."),
    ("Velocidad de transmisión", "kbps / Mbps", "Dimensionamiento de tráfico IoT", "Media", "Modelar uplink mínimo garantizado."),
    ("Hardware disponible", "categoría", "Factibilidad de despliegue", "Alta", "Clasificar módulos, terminales y madurez comercial."),
]

for values in rows:
    row = vars_table.add_row()
    for idx, value in enumerate(values):
        cell = row.cells[idx]
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(value)
        apply_font(r, 10)

closing = doc.add_paragraph()
closing.paragraph_format.space_before = Pt(10)
closing.paragraph_format.space_after = Pt(0)
cr1 = closing.add_run("Criterio de exito: ")
apply_font(cr1, 11, bold=True)
cr2 = closing.add_run(
    "el análisis debe permitir responder dónde, cuándo y con qué constelación existe cobertura IoT útil para operar en condiciones australes y polares."
)
apply_font(cr2, 11)

doc.save(OUTPUT_PATH)
print(OUTPUT_PATH)
