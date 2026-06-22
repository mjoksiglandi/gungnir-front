from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = r"C:\Users\juan.cornejo\Documents\gugnir v2\Terminos_de_Referencia_Practicante_Cobertura_IoT.docx"


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="DADCE0", size=8):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def apply_font(run, size, bold=False, color="000000", name="Arial"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.orientation = WD_ORIENTATION.PORTRAIT
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(11)

for style_name, size, color, before, after in (
    ("Heading 1", 20, "000000", 20, 6),
    ("Heading 2", 16, "000000", 18, 6),
    ("Heading 3", 14, "434343", 16, 4),
):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = False
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(3)
run = title.add_run("Términos de Referencia")
apply_font(run, 26)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(10)
subtitle_run = subtitle.add_run(
    "Practicante para investigación sobre cobertura satelital IoT en zonas australes y polares"
)
apply_font(subtitle_run, 11, color="555555")

meta = doc.add_table(rows=7, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
meta.autofit = False
meta.columns[0].width = Inches(2.0)
meta.columns[1].width = Inches(4.5)
set_table_borders(meta)

rows = [
    ("Cargo o encargo", "Practicante de investigación aplicada en conectividad satelital IoT."),
    ("Propósito", "Levantar, comparar y sintetizar información técnica sobre constelaciones satelitales con potencial uso en comunicaciones IoT para entornos australes y polares."),
    ("Unidad usuaria", "Equipo de análisis geoespacial, simulación operacional o innovación tecnológica."),
    ("Duración referencial", "4 a 8 semanas, según disponibilidad y profundidad del análisis."),
    ("Modalidad de trabajo", "Revisión documental, estructuración de base comparativa y desarrollo de productos cartográficos o analíticos."),
    ("Resultado esperado", "Informe base, matriz comparativa, mapas preliminares y recomendación técnica inicial."),
    ("Supervisor sugerido", "Responsable técnico del área de simulación, sistemas satelitales o planificación de capacidades."),
]

for row, (label, value) in zip(meta.rows, rows):
    for idx, text in enumerate((label, value)):
        cell = row.cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        apply_font(r, 10.5, bold=(idx == 0), color="000000" if idx == 0 else "111111")


def add_bullets(heading, items):
    doc.add_paragraph(heading, style="Heading 2")
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(item)
        apply_font(r, 11)


add_bullets(
    "1. Contexto del encargo",
    [
        "La organización requiere una línea base técnica para entender qué constelaciones satelitales podrían ofrecer cobertura útil para aplicaciones IoT en regiones australes, subantárticas y polares.",
        "El estudio debe servir como insumo inicial para decisiones de exploración tecnológica, simulación operativa y priorización de futuras pruebas o adquisiciones.",
    ],
)

add_bullets(
    "2. Objetivo general",
    [
        "Evaluar comparativamente constelaciones satelitales relevantes para comunicaciones IoT, con énfasis en cobertura, desempeño esperado y factibilidad de implementación en zonas australes y polares.",
    ],
)

add_bullets(
    "3. Objetivos específicos",
    [
        "Identificar constelaciones satelitales prioritarias para el análisis, tales como Iridium, Orbcomm, Swarm, Astrocast, Kinéis, Sateliot, Myriota u otras que el supervisor valide.",
        "Definir el área geográfica de estudio con criterio explícito, por ejemplo Magallanes, Paso Drake, Antártica chilena, mar austral o bandas latitudinales sobre 50°S y 60°S.",
        "Levantar variables comparables por constelación: frecuencias, número de satélites, cobertura declarada, footprint estimado, velocidades de transmisión, latencia referencial y disponibilidad de hardware compatible.",
        "Construir una matriz homogénea que permita comparar cobertura potencial, continuidad de servicio y madurez del ecosistema comercial.",
        "Generar mapas de calor, mapas de cobertura o productos equivalentes que faciliten la lectura espacial del análisis.",
    ],
)

add_bullets(
    "4. Alcance mínimo del trabajo",
    [
        "Revisión bibliográfica y documental basada en fuentes oficiales, literatura técnica, reguladores, operadores y proveedores de hardware.",
        "Comparación de al menos cinco constelaciones o soluciones satelitales relevantes.",
        "Elaboración de una base de datos estructurada con variables comparables y trazabilidad de fuentes.",
        "Producción de al menos tres visualizaciones geoespaciales o comparativas útiles para toma de decisión.",
        "Identificación de vacíos de información, supuestos metodológicos y riesgos de interpretación.",
    ],
)

add_bullets(
    "5. Preguntas que la investigación debe responder",
    [
        "Qué constelaciones muestran mejor potencial de cobertura para latitudes australes y polares.",
        "Qué soluciones parecen más adecuadas para sensores de baja tasa, telemetría intermitente o mensajería IoT.",
        "Qué bandas, velocidades y restricciones técnicas condicionan el uso real en terreno.",
        "Qué hardware comercial está disponible y qué tan viable es su integración en despliegues reales.",
        "Dónde existen mayores brechas de información y qué validaciones futuras serían necesarias.",
    ],
)

add_bullets(
    "6. Metodología esperada",
    [
        "Definir un marco de comparación con criterios homogéneos y una ficha estándar por constelación.",
        "Separar claramente datos confirmados, estimaciones, supuestos e inferencias del practicante.",
        "Utilizar herramientas geoespaciales, hojas de cálculo o scripts reproducibles cuando agreguen valor y trazabilidad.",
        "Incluir referencias completas de cada dato técnico relevante.",
    ],
)

doc.add_paragraph("7. Entregables", style="Heading 2")

deliverables = doc.add_table(rows=1, cols=4)
deliverables.alignment = WD_TABLE_ALIGNMENT.LEFT
deliverables.autofit = False
for width, col in zip((2.1, 1.1, 1.1, 2.2), deliverables.columns):
    col.width = Inches(width)
set_table_borders(deliverables)

headers = ["Entregable", "Formato", "Momento", "Contenido mínimo"]
for idx, text in enumerate(headers):
    cell = deliverables.rows[0].cells[idx]
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    r = p.add_run(text)
    apply_font(r, 10, bold=True)

deliverable_rows = [
    ("Plan de trabajo", "1-2 páginas", "Semana 1", "Objetivo, alcance, fuentes preliminares, cronograma y enfoque metodológico."),
    ("Matriz comparativa", "XLSX o CSV", "Semana 2-3", "Variables por constelación con trazabilidad de fuentes."),
    ("Mapas o visualizaciones", "PDF o imágenes", "Semana 3-4", "Cobertura o intensidad comparativa sobre áreas definidas."),
    ("Informe final", "DOCX o PDF", "Cierre", "Hallazgos, limitaciones, recomendación preliminar y próximos pasos."),
    ("Presentación breve", "PPTX o PDF", "Cierre", "Resumen ejecutivo de 10 a 15 minutos."),
]

for values in deliverable_rows:
    row = deliverables.add_row()
    for idx, value in enumerate(values):
        cell = row.cells[idx]
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(value)
        apply_font(r, 10)

add_bullets(
    "8. Criterios de evaluación del practicante",
    [
        "Claridad para acotar el problema y definir un método de trabajo realista.",
        "Rigor en el uso de fuentes y capacidad para distinguir información confirmada de supuestos.",
        "Calidad de la matriz comparativa y consistencia entre variables.",
        "Capacidad de traducir datos técnicos en visualizaciones útiles para la toma de decisión.",
        "Claridad de redacción, orden del informe y calidad de la recomendación final.",
    ],
)

add_bullets(
    "9. Resultado mínimo aceptable",
    [
        "Comparación documentada de al menos cinco constelaciones relevantes.",
        "Una base estructurada reutilizable por el equipo.",
        "Tres o más visualizaciones comprensibles y presentables.",
        "Una recomendación preliminar argumentada, aunque existan brechas de información.",
    ],
)

add_bullets(
    "10. Supuestos y limitaciones",
    [
        "El trabajo podrá combinar información pública, documentación comercial y estimaciones técnicas; no necesariamente validará desempeño real en terreno.",
        "Las cifras de cobertura o rendimiento deberán explicitar su fuente y el nivel de confianza asociado.",
        "Cualquier ausencia de datos deberá documentarse como hallazgo y no rellenarse sin justificación.",
    ],
)

closing = doc.add_paragraph()
closing.paragraph_format.space_before = Pt(10)
cr1 = closing.add_run("Nota para la convocatoria: ")
apply_font(cr1, 11, bold=True)
cr2 = closing.add_run(
    "si quieres atraer mejores postulantes, conviene adjuntar una lista inicial de constelaciones sugeridas, el área geográfica prioritaria y el formato esperado del producto final."
)
apply_font(cr2, 11)

doc.save(OUTPUT_PATH)
print(OUTPUT_PATH)
